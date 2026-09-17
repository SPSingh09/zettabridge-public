#!/usr/bin/env python3
"""Build Zerodha demo presentation from product docs + preliminary decks."""

from __future__ import annotations

import re
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    venv_dir = Path(__file__).resolve().parent / ".ppt-venv"
    venv_python = venv_dir / "bin" / "python"
    if sys.platform == "win32":
        venv_python = venv_dir / "Scripts" / "python.exe"
    if not venv_python.exists():
        subprocess.check_call([sys.executable, "-m", "venv", str(venv_dir)])
        subprocess.check_call([str(venv_python), "-m", "pip", "install", "python-pptx", "-q"])
    subprocess.check_call([str(venv_python), str(Path(__file__).resolve())] + sys.argv[1:])
    raise SystemExit(0)

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "presentations"
DIAGRAMS_DIR = OUT_DIR / "diagrams"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

# Brand colours
C_BG = RGBColor(15, 23, 42)       # slate-900
C_TITLE = RGBColor(248, 250, 252)
C_BODY = RGBColor(203, 213, 225)
C_ACCENT = RGBColor(59, 130, 246)  # blue-500
C_ACCENT2 = RGBColor(139, 92, 246) # violet-500
C_GREEN = RGBColor(74, 222, 128)
C_RED = RGBColor(248, 113, 113)
C_AMBER = RGBColor(251, 191, 36)
C_MUTED = RGBColor(148, 163, 184)
C_CARD = RGBColor(30, 41, 59)


def extract_pptx_text(path: Path) -> list[list[str]]:
    if not path.exists():
        return []
    slides: list[list[str]] = []
    with zipfile.ZipFile(path) as z:
        names = sorted(
            [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
            key=lambda x: int(re.search(r"slide(\d+)", x).group(1)),
        )
        for name in names:
            root = ET.fromstring(z.read(name))
            texts = [t.text.strip() for t in root.findall(".//a:t", NS) if t.text and t.text.strip()]
            slides.append(texts)
    return slides


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_footer(slide, text: str = "ZettaBridge — Zerodha Integration Demo | Confidential") -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED
    p.alignment = PP_ALIGN.CENTER


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    tbox = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.2))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(40)
    tp.font.bold = True
    tp.font.color.rgb = C_TITLE
    tp.alignment = PP_ALIGN.CENTER

    sbox = slide.shapes.add_textbox(Inches(1.0), Inches(3.3), Inches(11.0), Inches(1.5))
    sp = sbox.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(18)
    sp.font.color.rgb = C_BODY
    sp.alignment = PP_ALIGN.CENTER

    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(4.2), Inches(5.0), Inches(4.6), Inches(0.55))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    badge.line.color.rgb = C_ACCENT
    btf = badge.text_frame
    btf.text = "July 2026 · Closed Beta · Paper + Kite Publisher"
    btf.paragraphs[0].font.size = Pt(12)
    btf.paragraphs[0].font.color.rgb = C_TITLE
    btf.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_footer(slide)


def add_section_slide(prs: Presentation, section: str, subtitle: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(3.1), Inches(13.33), Inches(1.4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(30, 58, 138)
    bar.line.fill.background()
    tbox = slide.shapes.add_textbox(Inches(0.8), Inches(3.25), Inches(11.5), Inches(1.0))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = section
    tp.font.size = Pt(32)
    tp.font.bold = True
    tp.font.color.rgb = C_TITLE
    tp.alignment = PP_ALIGN.CENTER
    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10.8), Inches(0.8))
        sp = sbox.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(16)
        sp.font.color.rgb = C_BODY
        sp.alignment = PP_ALIGN.CENTER
    add_footer(slide)


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str], note: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    h = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12.0), Inches(0.7))
    hp = h.text_frame.paragraphs[0]
    hp.text = title
    hp.font.size = Pt(28)
    hp.font.bold = True
    hp.font.color.rgb = C_TITLE

    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.35), Inches(11.8), Inches(5.3))
    tf = body.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(17)
        p.font.color.rgb = C_BODY
        p.space_after = Pt(10)

    if note:
        nbox = slide.shapes.add_textbox(Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.4))
        np = nbox.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(11)
        np.font.italic = True
        np.font.color.rgb = C_MUTED
    add_footer(slide)


def add_two_column_slide(prs: Presentation, title: str, left_title: str, left_items: list[str], right_title: str, right_items: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    h = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12.0), Inches(0.7))
    hp = h.text_frame.paragraphs[0]
    hp.text = title
    hp.font.size = Pt(26)
    hp.font.bold = True
    hp.font.color.rgb = C_TITLE

    for col, (ctitle, items, x, accent) in enumerate(
        [
            (left_title, left_items, 0.7, C_GREEN),
            (right_title, right_items, 6.9, C_RED),
        ]
    ):
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.3), Inches(5.9), Inches(5.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = accent
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)
        p0 = tf.paragraphs[0]
        p0.text = ctitle
        p0.font.size = Pt(18)
        p0.font.bold = True
        p0.font.color.rgb = accent
        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(14)
            p.font.color.rgb = C_BODY
            p.space_after = Pt(6)
    add_footer(slide)


def add_table_slide(prs: Presentation, title: str, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    h = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12.0), Inches(0.7))
    hp = h.text_frame.paragraphs[0]
    hp.text = title
    hp.font.size = Pt(24)
    hp.font.bold = True
    hp.font.color.rgb = C_TITLE

    nrows = len(rows) + 1
    ncols = len(headers)
    table = slide.shapes.add_table(nrows, ncols, Inches(0.6), Inches(1.25), Inches(12.1), Inches(0.45 * nrows)).table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)

    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 58, 138)
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(12)
            p.font.color.rgb = C_TITLE

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD if r % 2 else RGBColor(24, 33, 51)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = C_BODY
    add_footer(slide)


def _box(slide, x, y, w, h, text, fill, line, size=11, bold=False):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = C_TITLE
    p.alignment = PP_ALIGN.CENTER
    return shape


def _arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = C_MUTED
    conn.line.width = Pt(1.5)


def add_flow_diagram_slide(prs: Presentation, title: str, nodes: list[tuple[str, float, float, float, float]], arrows: list[tuple[float, float, float, float]], caption: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    h = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12.0), Inches(0.6))
    hp = h.text_frame.paragraphs[0]
    hp.text = title
    hp.font.size = Pt(24)
    hp.font.bold = True
    hp.font.color.rgb = C_TITLE

    for text, x, y, w, hgt in nodes:
        _box(slide, x, y, w, hgt, text, C_CARD, C_ACCENT, size=10)
    for x1, y1, x2, y2 in arrows:
        _arrow(slide, x1, y1, x2, y2)

    if caption:
        cbox = slide.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.35))
        cp = cbox.text_frame.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(10)
        cp.font.italic = True
        cp.font.color.rgb = C_MUTED
        cp.alignment = PP_ALIGN.CENTER
    add_footer(slide)


def add_sequence_slide(prs: Presentation) -> None:
    steps = [
        ("1", "Signal source", "POST JSON to webhook URL"),
        ("2", "ZettaBridge", "Validate → HTTP 202 + request_id"),
        ("3", "ZettaBridge", "Queue job · status → queued"),
        ("4", "ZettaBridge", "Publisher handoff · pending_confirmation"),
        ("5", "User", "Open trade in ZettaBridge dashboard"),
        ("6", "User", "Review & confirm order on Zerodha Kite"),
        ("7", "ZettaBridge", "Audit: submitted / filled / rejected"),
    ]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    h = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12.0), Inches(0.6))
    hp = h.text_frame.paragraphs[0]
    hp.text = "Kite Publisher — end-to-end confirmation flow"
    hp.font.size = Pt(24)
    hp.font.bold = True
    hp.font.color.rgb = C_TITLE

    y = 1.1
    for num, who, action in steps:
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.8), Inches(y), Inches(0.35), Inches(0.35))
        circ.fill.solid()
        circ.fill.fore_color.rgb = RGBColor(30, 58, 138)
        circ.line.color.rgb = C_ACCENT
        ctf = circ.text_frame
        ctf.text = num
        ctf.paragraphs[0].font.size = Pt(11)
        ctf.paragraphs[0].font.color.rgb = C_TITLE
        ctf.paragraphs[0].alignment = PP_ALIGN.CENTER

        tbox = slide.shapes.add_textbox(Inches(1.35), Inches(y - 0.02), Inches(11.0), Inches(0.55))
        tp = tbox.text_frame.paragraphs[0]
        tp.text = f"{who} — {action}"
        tp.font.size = Pt(14)
        tp.font.color.rgb = C_BODY if who != "User" else C_AMBER
        if who in ("ZettaBridge", "User"):
            tp.runs[0].font.bold = True
        y += 0.72

    note = slide.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.5))
    np = note.text_frame.paragraphs[0]
    np.text = "ZettaBridge never stores Zerodha login passwords. Live orders require explicit user confirmation on Zerodha."
    np.font.size = Pt(11)
    np.font.italic = True
    np.font.color.rgb = C_MUTED
    add_footer(slide)


def add_architecture_slide(prs: Presentation) -> None:
    layers = [
        ("Signal Source", "TradingView · AI Bot · Python", 0.9),
        ("Cloudflare Edge", "TLS · DDoS · WAF", 1.55),
        ("API Gateway", "Token auth · rate limits · validation", 2.2),
        ("Execution Engine", "Dedup · guards · workflow routing", 2.85),
        ("Redis", "Dedup window · session state", 3.5),
        ("Broker Workflow", "Paper engine · Kite Publisher handoff", 4.15),
        ("Market Access", "Paper sim · Zerodha (user-confirmed)", 4.8),
        ("Audit DB", "PostgreSQL — full event log", 5.45),
    ]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, C_BG)
    h = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12.0), Inches(0.6))
    hp = h.text_frame.paragraphs[0]
    hp.text = "ZettaBridge execution architecture"
    hp.font.size = Pt(24)
    hp.font.bold = True
    hp.font.color.rgb = C_TITLE

    for label, sub, y in layers:
        _box(slide, 2.5, y, 8.3, 0.5, f"{label}\n{sub}", C_CARD, C_ACCENT2 if "Engine" in label else C_ACCENT, size=10)
        if y < 5.45:
            _arrow(slide, 6.4, y + 0.5, 6.4, y + 0.58)

    zbadge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.5), Inches(1.5), Inches(2.2))
    zbadge.fill.solid()
    zbadge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    zbadge.line.color.rgb = C_ACCENT2
    ztf = zbadge.text_frame
    ztf.text = "ZettaBridge\nExecution\nPlatform"
    for p in ztf.paragraphs:
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_TITLE
        p.alignment = PP_ALIGN.CENTER
    add_footer(slide)


def add_security_layers_slide(prs: Presentation) -> None:
    layers = [
        ("1. Edge", "Cloudflare — TLS termination, DDoS protection, WAF"),
        ("2. Ingest", "Hashed webhook tokens, rate limits (60/min), dedup (2s), trading-hour guards"),
        ("3. Validation", "action · symbol · comment · plan gates — reject before queue where possible"),
        ("4. Processing", "Async workers — no synchronous broker call in webhook response"),
        ("5. Publisher", "User must confirm on Zerodha — no auto-placement without confirmation"),
        ("6. Audit", "request_id → ingest log → trade record → dashboard — full traceability"),
        ("7. Storage", "AES-256-GCM for secrets · PostgreSQL per-account isolation"),
    ]
    add_bullet_slide(prs, "Security controls at every stage", [f"{a}: {b}" for a, b in layers],
                    note="Internal adapter topology uses private subnets, service tokens, and NAT egress for Kite API (staging).")


def build_presentation(prep_slides: list[list[str]], demo_slides: list[list[str]]) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "ZettaBridge × Zerodha",
        "Product capability, safety, security & transparency\nDemonstrating our platform for Kite Connect / Personal API consideration",
    )

    add_section_slide(prs, "Purpose of this demo", "Requesting Zerodha's review of our closed-beta platform")

    add_bullet_slide(
        prs,
        "Why we are here today",
        [
            "Demonstrate what ZettaBridge is — an execution router for TradingView, bots, and custom signals",
            "Show end-to-end transparency: signal ingest → validation → paper or Publisher workflow → audit",
            "Prove safety & security posture: what we store, what we never store, controls at each layer",
            "Explain current Kite Publisher integration (closed beta) and its deliberate limitations",
            "Present our readiness for Kite Connect / Kite Personal API — with SEBI Algo-ID compliance built in",
            "Seek Zerodha's permission to proceed with approved API integration for eligible users",
        ],
    )

    add_section_slide(prs, "What is ZettaBridge?")

    add_bullet_slide(
        prs,
        "ZettaBridge in one slide",
        [
            "Webhook-based execution platform for Indian retail algo traders",
            "Receives signals from TradingView alerts, Python scripts, AI bots, or any HTTPS client",
            "Routes each signal to a configured workflow: built-in paper trading or Zerodha live (Publisher today)",
            "Returns HTTP 202 Accepted — signal is validated and queued, not executed synchronously in the response",
            "Maintains a complete audit trail: request_id, ingest logs, trade status, execution logs",
            "We are an execution router — not a broker, not an investment adviser, not a custodian",
        ],
        note="Closed beta: Paper trading available · Zerodha Kite Publisher in closed beta · Kite Connect not yet public",
    )

    add_two_column_slide(
        prs,
        "Compliance role boundary",
        "ZettaBridge does",
        [
            "Validate & route webhook signals",
            "Enforce plan limits, dedup, trading-hour guards",
            "Log full audit trail per account",
            "Attach SEBI Algo-ID when configured (Kite Connect path)",
            "Hand off Publisher orders for user confirmation",
        ],
        "ZettaBridge does NOT",
        [
            "Hold customer funds or securities",
            "Store Zerodha login passwords",
            "Provide investment advice",
            "Execute discretionary trades without user action",
            "Bypass Zerodha's order confirmation for Publisher flow",
        ],
    )

    add_section_slide(prs, "Architecture", "Enterprise-grade execution stack")

    add_architecture_slide(prs)

    add_flow_diagram_slide(
        prs,
        "Webhook ingest sequence",
        [
            ("Signal source\nTradingView · Python · AI", 4.8, 1.0, 3.7, 0.55),
            ("Token lookup + validation\naction · comment · hours · rate", 4.8, 1.85, 3.7, 0.65),
            ("HTTP response\n202 / 4xx / 409 / 429", 4.8, 2.8, 3.7, 0.55),
            ("Async worker\nsymbol · qty · routing", 4.8, 3.65, 3.7, 0.55),
            ("Outcome\nqueued → filled / rejected /\npending_confirmation", 4.8, 4.5, 3.7, 0.7),
        ],
        [
            (6.65, 1.55, 6.65, 1.85),
            (6.65, 2.5, 6.65, 2.8),
            (6.65, 3.35, 6.65, 3.65),
            (6.65, 4.2, 6.65, 4.5),
        ],
        caption="Synchronous validation at ingest · async processing after HTTP 202 · no broker redirect in webhook response",
    )

    add_section_slide(prs, "End-to-end flows", "Paper trading and Kite Publisher")

    add_flow_diagram_slide(
        prs,
        "Paper vs Kite Publisher — same ingest, different destination",
        [
            ("POST /v1/webhook/{token}\n202 + request_id", 4.6, 0.95, 4.1, 0.55),
            ("Route by webhook\nconfiguration", 5.0, 1.75, 3.3, 0.5),
            ("Paper engine\nSimulated balance", 1.0, 2.85, 3.0, 0.55),
            ("Paper fill\nfilled / rejected", 1.0, 3.65, 3.0, 0.5),
            ("Validate & queue", 9.3, 2.85, 3.0, 0.5),
            ("pending_confirmation", 9.3, 3.55, 3.0, 0.45),
            ("User confirms\non Zerodha Kite", 9.3, 4.25, 3.0, 0.55),
            ("submitted / filled / rejected", 9.3, 5.05, 3.0, 0.5),
        ],
        [
            (6.15, 1.5, 6.15, 1.75),
            (5.3, 2.2, 2.5, 2.85),
            (7.0, 2.2, 10.8, 2.85),
            (2.5, 3.4, 2.5, 3.65),
            (10.8, 3.35, 10.8, 3.55),
            (10.8, 4.0, 10.8, 4.25),
            (10.8, 4.8, 10.8, 5.05),
        ],
        caption="Destination is set per webhook — not per signal payload. Publisher requires explicit user confirmation.",
    )

    add_bullet_slide(
        prs,
        "Paper trading workflow",
        [
            "User creates a paper account — no broker credentials required",
            "Webhook links to paper account; signals simulate orders against configured defaults",
            "Outcomes: filled or rejected — resolved entirely inside ZettaBridge",
            "Used for strategy validation before enabling any live Zerodha workflow",
            "Same ingest path, guards, dedup, and audit logging as live webhooks",
        ],
    )

    add_sequence_slide(prs)

    add_table_slide(
        prs,
        "Trade status lifecycle",
        ["Status", "Meaning", "Typical workflow"],
        [
            ["queued", "Signal accepted & queued", "All webhooks after 202"],
            ["pending_confirmation", "Awaiting user on Zerodha", "Publisher only"],
            ["submitted", "Order submission recorded", "After Zerodha confirmation"],
            ["filled", "Paper fill or broker fill", "Terminal success"],
            ["rejected", "Validation or workflow error", "May occur after 202"],
            ["cancelled", "User or system cancelled", "Either workflow"],
        ],
        col_widths=[2.2, 4.5, 5.4],
    )

    add_section_slide(prs, "Data & security", "Transparency on what we store")

    add_two_column_slide(
        prs,
        "Data boundary — what ZettaBridge stores",
        "We store (operational necessity)",
        [
            "Webhook configuration & defaults",
            "Hashed webhook tokens (never plaintext)",
            "Paper account records",
            "Broker configuration labels (not login passwords)",
            "Order attempts & full audit logs",
            "Optional order tags (e.g. SEBI Algo-ID)",
        ],
        "We never store",
        [
            "Zerodha login password",
            "Customer funds or securities",
            "TradingView credentials",
            "Bank or payment credentials",
            "Discretionary trading instructions beyond user config",
        ],
    )

    add_security_layers_slide(prs)

    add_flow_diagram_slide(
        prs,
        "Audit trail — end-to-end traceability",
        [
            ("202 response\nrequest_id", 1.2, 2.5, 2.4, 0.55),
            ("Ingest log\naccepted / rejected / dedup", 4.0, 2.5, 2.6, 0.55),
            ("Trade record\nstatus · error_code", 6.9, 2.5, 2.5, 0.55),
            ("Dashboard\nWebhook & Execution Logs", 9.7, 2.5, 2.5, 0.55),
        ],
        [
            (3.6, 2.78, 4.0, 2.78),
            (6.6, 2.78, 6.9, 2.78),
            (9.4, 2.78, 9.7, 2.78),
        ],
        caption="Every signal traceable from source request_id through to final trade outcome",
    )

    add_section_slide(prs, "Kite Publisher today", "What works · what is intentionally limited")

    add_table_slide(
        prs,
        "Kite Publisher — current closed-beta capabilities",
        ["Capability", "Publisher (today)", "Notes"],
        [
            ["Signal ingest", "HTTP 202 — async queue", "No synchronous order in response"],
            ["Order placement", "User confirms on Zerodha Kite", "Manual confirmation required"],
            ["Actions", "BUY / SELL", "CLOSE not supported on Publisher webhooks"],
            ["OAuth session", "Not required for Publisher handoff", "No stored Zerodha password"],
            ["Audit", "Full request_id trail", "Dashboard + PostgreSQL logs"],
            ["Algo-ID", "Not required (manual flow)", "Publisher is user-confirmed, not algo-automated"],
        ],
        col_widths=[3.0, 4.5, 4.6],
    )

    add_bullet_slide(
        prs,
        "Kite Publisher — deliberate limitations",
        [
            "User must manually confirm every order on Zerodha — cannot fully automate from TradingView webhook alone",
            "Browser-redirect handoff — not suitable for unattended server-side algo execution",
            "No programmatic order status stream equivalent to Kite Connect WebSocket",
            "Daily session constraints — Publisher relies on user being available to confirm",
            "Limited to BUY/SELL actions at webhook configuration level",
            "We chose Publisher for closed beta to demonstrate safety without requesting Connect prematurely",
        ],
        note="These limitations are features of the Publisher model — not bugs. They protect users and demonstrate our conservative approach.",
    )

    add_section_slide(prs, "Kite Connect / Personal API", "What integration would unlock")

    add_table_slide(
        prs,
        "Publisher vs Kite Connect — honest comparison",
        ["Dimension", "Kite Publisher (today)", "Kite Connect / Personal API (requested)"],
        [
            ["Automation", "User confirms each order on Kite", "Programmatic place/modify/cancel via API"],
            ["Latency", "Human confirmation step", "Sub-second API routing after guards"],
            ["Webhook → order", "Signal queued; user acts later", "Signal → validated → API order (with Algo-ID)"],
            ["Order status", "Poll + callback after confirmation", "Real-time order updates via API/WS"],
            ["SEBI compliance", "Manual flow — no Algo-ID needed", "Algo-ID registration & per-order tagging required"],
            ["User safety", "Maximum — every order reviewed", "Guards + Algo-ID + rate limits + audit"],
            ["Our readiness", "Shipped in closed beta", "Architecture built; awaiting Zerodha approval"],
        ],
        col_widths=[2.5, 4.8, 4.8],
    )

    add_bullet_slide(
        prs,
        "Why we are ready for Kite Connect approval",
        [
            "Execution architecture already separates Core (ingest, guards, audit) from Zerodha adapter (Kite API, OAuth, NAT egress)",
            "SEBI Algo-ID tagging implemented — attached per order when configured on broker credential",
            "Encrypted credential storage (AES-256-GCM) — access tokens in PostgreSQL, never logged",
            "OAuth callback on dedicated adapter host — Core never receives Kite request_token on public internet",
            "Plan gates, dedup (2s), rate limits (60/min), trading-hour guards — all enforced before API call",
            "Full audit trail already production-grade — Connect adds API order IDs to existing log chain",
            "Staging deployed on AWS ap-south-1 with private subnets, static NAT IP for Kite whitelist",
        ],
    )

    add_flow_diagram_slide(
        prs,
        "Proposed Kite Connect flow (post-approval)",
        [
            ("TradingView / Bot\nPOST webhook", 0.8, 2.2, 2.5, 0.55),
            ("ZettaBridge Core\nvalidate · dedup · guards", 3.7, 2.2, 2.8, 0.55),
            ("Zerodha Adapter\nOAuth session · PlaceOrder", 7.0, 2.2, 2.8, 0.55),
            ("Kite Connect API\norder + Algo-ID tag", 10.1, 2.2, 2.5, 0.55),
            ("Audit DB\norder_id · status · Algo-ID", 5.5, 4.0, 3.5, 0.55),
        ],
        [
            (3.3, 2.48, 3.7, 2.48),
            (6.5, 2.48, 7.0, 2.48),
            (9.8, 2.48, 10.1, 2.48),
            (8.4, 2.75, 7.25, 4.0),
            (6.75, 2.75, 6.75, 4.0),
        ],
        caption="Same ingest safety layer — Connect replaces Publisher handoff with approved API execution under Algo-ID compliance",
    )

    add_section_slide(prs, "Operational controls")

    add_table_slide(
        prs,
        "Rate limits, dedup & plan gates",
        ["Control", "Default", "Purpose"],
        [
            ["Per-webhook rate limit", "60 requests / minute", "Prevent signal flooding"],
            ["Dedup window", "2 seconds", "Suppress duplicate alerts"],
            ["Free plan daily cap", "100 ingests / day", "Abuse prevention"],
            ["Broker order rate", "10 orders / sec per config", "Exchange/broker protection"],
            ["Plan gate (live)", "Pro / Pro Plus + beta approval", "Controlled live rollout"],
            ["HTTP 409 on dedup", "No retry needed", "Original signal still processing"],
        ],
        col_widths=[3.5, 3.5, 5.1],
    )

    add_section_slide(prs, "What we are asking Zerodha")

    add_bullet_slide(
        prs,
        "Our request",
        [
            "Review this demonstration of ZettaBridge's current closed-beta platform",
            "Acknowledge our safety, security, and audit posture demonstrated via Kite Publisher integration",
            "Grant permission to proceed with Kite Connect / Kite Personal API integration for approved users",
            "Whitelist our staging/production NAT egress IP for Kite API access",
            "Guide us on SEBI Algo-ID registration requirements for our platform category",
            "We commit to: no public Connect promises until Zerodha-approved · conservative rollout · full audit transparency",
        ],
    )

    add_bullet_slide(
        prs,
        "Next steps after approval",
        [
            "Enable Kite Connect OAuth for closed-beta users on Pro / Pro Plus plans",
            "Attach registered Algo-ID to every programmatic order",
            "Maintain Publisher as an option for users who prefer manual confirmation",
            "Publish integration documentation and error-code reference for transparency",
            "Ongoing audit log access for compliance review on request",
            "Contact: team@zettabridge.com · staging: api.staging.zettabridge.net",
        ],
    )

    add_title_slide(
        prs,
        "Thank you",
        "ZettaBridge — Safe, transparent execution routing for Indian algo traders\n\nWe welcome your questions and feedback.",
    )

    # Optional: log extracted preliminary content for traceability
    trace_path = OUT_DIR / "preliminary_deck_notes.txt"
    lines = ["# Extracted from preliminary decks\n"]
    for label, slides in [("Preparation", prep_slides), ("Broker Demo", demo_slides)]:
        lines.append(f"\n## {label}\n")
        if not slides:
            lines.append("(file not found or empty)\n")
            continue
        for i, slide_texts in enumerate(slides, 1):
            lines.append(f"### Slide {i}\n")
            lines.extend(f"- {t}" for t in slide_texts)
            lines.append("")
    trace_path.write_text("\n".join(lines), encoding="utf-8")

    out_path = OUT_DIR / "ZettaBridge_Zerodha_Demo_Final.pptx"
    prs.save(str(out_path))
    return out_path


def find_preliminary_decks() -> tuple[Path, Path]:
    candidates = [
        Path(r"C:\Users\Surya\Downloads"),
        Path("/mnt/c/Users/Surya/Downloads"),
    ]
    prep_name = "ZettaBridge_Zerodha_Demo_Preparation.pptx"
    demo_name = "ZettaBridge_Zerodha_Broker_Demo.pptx"
    prep = demo = None
    for base in candidates:
        p = base / prep_name
        d = base / demo_name
        if p.exists():
            prep = p
        if d.exists():
            demo = d
    return prep or Path(prep_name), demo or Path(demo_name)


def main() -> None:
    prep_path, demo_path = find_preliminary_decks()
    print(f"Reading: {prep_path} ({'found' if prep_path.exists() else 'missing'})")
    print(f"Reading: {demo_path} ({'found' if demo_path.exists() else 'missing'})")
    prep_slides = extract_pptx_text(prep_path)
    demo_slides = extract_pptx_text(demo_path)
    out = build_presentation(prep_slides, demo_slides)
    print(f"Created: {out}")
    # Also copy to Downloads for easy access (WSL + Windows)
    for dl in [
        Path("/mnt/c/Users/Surya/Downloads") / out.name,
        Path(r"C:\Users\Surya\Downloads") / out.name,
    ]:
        if dl.parent.exists():
            import shutil
            shutil.copy2(out, dl)
            print(f"Copied to: {dl}")
            break


if __name__ == "__main__":
    main()
